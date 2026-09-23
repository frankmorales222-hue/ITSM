from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "word"
OUT.mkdir(parents=True, exist_ok=True)

NAVY = "0B2545"
BLUE = "2E74B5"
MUTED = "5B6B7A"
PALE = "E8EEF5"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None: node = OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def add_page_number(paragraph):
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instr); run._r.append(fldChar2)

def setup_doc(title, audience):
    d = Document(); sec = d.sections[0]
    sec.top_margin = Inches(0.8); sec.bottom_margin = Inches(0.75); sec.left_margin = Inches(0.85); sec.right_margin = Inches(0.85)
    sec.header_distance = Inches(0.35); sec.footer_distance = Inches(0.35)
    styles = d.styles
    normal = styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(10.5); normal.font.color.rgb = RGBColor.from_string(NAVY)
    normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.18
    for name, size, color, before, after in [('Heading 1',16,BLUE,16,8),('Heading 2',13,BLUE,12,6),('Heading 3',11.5,NAVY,9,4)]:
        s=styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    hp = sec.header.paragraphs[0]; hp.text = 'NORTHSTAR DESK  |  ' + audience.upper(); hp.style='Caption'; hp.runs[0].font.color.rgb=RGBColor.from_string(MUTED); hp.runs[0].font.bold=True
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; r=fp.add_run('Northstar Desk  •  '); r.font.color.rgb=RGBColor.from_string(MUTED); add_page_number(fp)
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(2); r=p.add_run('NORTHSTAR DESK'); r.font.name='Calibri'; r.font.size=Pt(10); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(BLUE)
    p=d.add_paragraph(); p.style='Title'; p.paragraph_format.space_after=Pt(4); r=p.add_run(title); r.font.name='Calibri'; r.font.size=Pt(27); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
    p=d.add_paragraph(audience); p.paragraph_format.space_after=Pt(14); p.runs[0].font.size=Pt(13); p.runs[0].font.color.rgb=RGBColor.from_string(MUTED)
    t=d.add_table(rows=1, cols=1); t.autofit=False; t.columns[0].width=Inches(6.65); c=t.cell(0,0); shade(c, PALE); set_cell_margins(c,140,160,140,160); rr=c.paragraphs[0].add_run('Quick reference'); rr.bold=True; rr.font.color.rgb=RGBColor.from_string(BLUE); c.add_paragraph('Use this manual with the current Northstar Desk interface. Labels may vary slightly by role and enabled modules.')
    d.add_paragraph()
    return d

def add_md(d, path):
    lines=Path(path).read_text(encoding='utf-8').splitlines(); first=True
    for line in lines:
        if not line.strip(): continue
        if first and line.startswith('# '): first=False; continue
        if line.startswith('### '): d.add_heading(line[4:], level=3)
        elif line.startswith('## '): d.add_heading(line[3:], level=1)
        elif line.startswith('# '): d.add_heading(line[2:], level=1)
        elif line.startswith('- '): d.add_paragraph(line[2:], style='List Bullet')
        elif line[:2].isdigit() and line[2:4] == '. ': d.add_paragraph(line[4:], style='List Number')
        elif line.startswith('1. '): d.add_paragraph(line[3:], style='List Number')
        else: d.add_paragraph(line)

def save(src, title, audience, filename):
    d=setup_doc(title,audience); add_md(d, src); d.save(OUT/filename)

save(ROOT/'docs/TECHNICIAN_GUIDE.md','Technician Manual','Daily ticket handling, communication, routing, and support operations','Northstar_Desk_Technician_Manual.docx')
save(ROOT/'docs/ADMIN_GUIDE.md','Administrator Manual','Configuration, security, integrations, reporting, and recovery','Northstar_Desk_Administrator_Manual.docx')
save(ROOT/'docs/END_USER_GUIDE.md','Requester Manual','How employees submit requests, follow updates, and get help','Northstar_Desk_Requester_Manual.docx')
print('Created 3 manuals in', OUT)
